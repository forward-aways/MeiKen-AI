"""Excel 解析：.xlsx/.xlsm 用 openpyxl，.xls 旧格式用 xlrd。

每个工作表输出为二级标题 + Markdown 表格。
"""
import xlrd
from openpyxl import load_workbook

from backend.documents.parsers.base import ParsedDocument, Parser, clip_text, rows_to_markdown

_MAX_ROWS = 500


class XlsxParser(Parser):
    extensions = (".xlsx", ".xlsm")

    def parse(self, path: str, filename: str) -> ParsedDocument:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            parts = []
            for ws in wb.worksheets:
                parts.append(f"## 工作表：{ws.title}")
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > _MAX_ROWS:
                        break
                    rows.append(list(row))
                parts.append(rows_to_markdown(rows))
            return ParsedDocument(
                text=clip_text("\n\n".join(parts)),
                meta={"sheets": len(wb.sheetnames)},
            )
        finally:
            wb.close()


class XlsParser(Parser):
    extensions = (".xls",)

    def parse(self, path: str, filename: str) -> ParsedDocument:
        book = xlrd.open_workbook(path)
        parts = []
        for sh in book.sheets():
            parts.append(f"## 工作表：{sh.name}")
            rows = [
                [sh.cell_value(r, c) for c in range(sh.ncols)]
                for r in range(min(sh.nrows, _MAX_ROWS + 1))
            ]
            parts.append(rows_to_markdown(rows))
        return ParsedDocument(
            text=clip_text("\n\n".join(parts)),
            meta={"sheets": book.nsheets},
        )
