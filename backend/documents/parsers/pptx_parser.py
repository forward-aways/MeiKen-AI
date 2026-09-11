"""PowerPoint（.pptx）解析：按幻灯片输出标题、文本与表格，备注以引用块附加。"""
from pptx import Presentation

from backend.documents.parsers.base import ParsedDocument, Parser, clip_text, rows_to_markdown


class PptxParser(Parser):
    extensions = (".pptx",)

    def parse(self, path: str, filename: str) -> ParsedDocument:
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            lines = [f"## 幻灯片 {i}"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs).strip()
                        if t:
                            lines.append(t)
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    lines.append(rows_to_markdown(rows))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"> 备注：{notes}")
            parts.append("\n\n".join(lines))
        return ParsedDocument(
            text=clip_text("\n\n---\n\n".join(parts)),
            meta={"slides": len(prs.slides)},
        )
