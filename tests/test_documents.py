"""文档解析器测试：验证各格式解析为 Markdown 文本。

运行：uv run --with pytest pytest tests/test_documents.py -v
"""
import pytest

from backend.documents import UnsupportedFormatError, extract_document


def test_text(tmp_path):
    p = tmp_path / "note.md"
    p.write_text("# 标题\n\n正文内容", encoding="utf-8")
    doc = extract_document(str(p), "note.md")
    assert "# 标题" in doc.text and "正文内容" in doc.text


def test_csv(tmp_path):
    p = tmp_path / "data.csv"
    p.write_text("名称,数量\n苹果,3\n香蕉,5\n", encoding="utf-8")
    doc = extract_document(str(p), "data.csv")
    assert "| 名称 | 数量 |" in doc.text
    assert "| 苹果 | 3 |" in doc.text
    assert doc.meta["rows"] == 2


def test_docx(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading("一级标题", level=1)
    doc.add_paragraph("正文段落")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列A"
    table.cell(0, 1).text = "列B"
    table.cell(1, 0).text = "值1"
    table.cell(1, 1).text = "值2"
    p = tmp_path / "doc.docx"
    doc.save(str(p))

    parsed = extract_document(str(p), "doc.docx")
    assert "# 一级标题" in parsed.text
    assert "正文段落" in parsed.text
    assert "| 列A | 列B |" in parsed.text
    assert parsed.meta["tables"] == 1


def test_xlsx(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", "数量"])
    ws.append(["苹果", 3])
    p = tmp_path / "book.xlsx"
    wb.save(str(p))

    parsed = extract_document(str(p), "book.xlsx")
    assert "## 工作表：数据" in parsed.text
    assert "| 名称 | 数量 |" in parsed.text
    assert parsed.meta["sheets"] == 1


def test_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "封面标题"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "要点一"
    p = tmp_path / "deck.pptx"
    prs.save(str(p))

    parsed = extract_document(str(p), "deck.pptx")
    assert "## 幻灯片 1" in parsed.text
    assert "封面标题" in parsed.text
    assert "要点一" in parsed.text
    assert parsed.meta["slides"] == 1


def test_html(tmp_path):
    html = (
        "<html><head><title>页面标题</title></head><body>"
        "<h1>大标题</h1><p>段落文字</p><ul><li>项目一</li></ul>"
        "<table><tr><th>甲</th><th>乙</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "</body></html>"
    )
    p = tmp_path / "page.html"
    p.write_text(html, encoding="utf-8")
    parsed = extract_document(str(p), "page.html")
    assert "# 页面标题" in parsed.text
    assert "# 大标题" in parsed.text
    assert "段落文字" in parsed.text
    assert "- 项目一" in parsed.text
    assert "| 甲 | 乙 |" in parsed.text


def test_pdf(tmp_path):
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    p = tmp_path / "empty.pdf"
    with open(str(p), "wb") as f:
        w.write(f)
    parsed = extract_document(str(p), "empty.pdf")
    assert parsed.meta["pages"] == 1


def test_unsupported(tmp_path):
    p = tmp_path / "file.xyz"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFormatError):
        extract_document(str(p), "file.xyz")
